"""Tests for document extraction and the DocumentAgent (Phase 14)."""

from __future__ import annotations

import io

from app.agents.documents import DocumentAgent
from app.services.documents import extract_text

# --- extraction --------------------------------------------------------------


def _make_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_docx() -> None:
    data = _make_docx("Invoice total 100")
    text = extract_text(data, mime_type=None, filename="doc.docx")
    assert "Invoice total 100" in text


def test_extract_xlsx() -> None:
    data = _make_xlsx([["Concept", "Amount"], ["Fee", "9.99"]])
    text = extract_text(data, mime_type=None, filename="sheet.xlsx")
    assert "Amount" in text
    assert "9.99" in text


def test_extract_corrupt_returns_empty() -> None:
    # Corrupt bytes must degrade to empty, never raise.
    assert extract_text(b"not a real pdf", mime_type="application/pdf") == ""


def test_extract_unknown_type_returns_empty() -> None:
    assert extract_text(b"whatever", mime_type="text/plain", filename="x.txt") == ""


# --- DocumentAgent -----------------------------------------------------------


async def test_agent_detects_invoice() -> None:
    result = DocumentAgent().analyze_document("This invoice is due. Amount due: 100")
    finding = result.findings[0]
    assert finding.kind == "invoice"
    assert finding.needs_review is True


async def test_agent_detects_contract() -> None:
    result = DocumentAgent().analyze_document("This agreement contains terms and conditions")
    assert result.findings[0].kind == "contract"


async def test_agent_empty_text() -> None:
    result = DocumentAgent().analyze_document("")
    assert not result.has_findings
    assert result.confidence == 0.0


async def test_agent_masks_numbers_in_summary() -> None:
    # A long account-like number must be masked in any surfaced text.
    result = DocumentAgent().analyze_document("Statement for account 4111111111111234 balance")
    assert result.has_findings
    assert "4111111111111234" not in result.findings[0].summary


async def test_agent_unknown_document() -> None:
    result = DocumentAgent().analyze_document("Just some random text with no keywords")
    assert not result.has_findings
