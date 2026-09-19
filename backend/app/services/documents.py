"""HELIOS — Document text extraction (Phase 14).

Extracts plain text from common document types (PDF, Excel, Word) so the
DocumentAgent can analyze attachments. Extraction is:

- **Text-only.** We pull text, never execute macros or embedded content. The raw
  bytes are treated as UNTRUSTED data.
- **Defensive.** Any parsing error returns empty text rather than raising — a
  corrupt or malicious file must never break the pipeline.
- **Lazy-imported.** The parser libraries are imported inside the functions so
  they are only needed when a document of that type is actually processed.

OCR for images is intentionally out of scope for this phase (documented as a
future enhancement).
"""

from __future__ import annotations

import io

from app.observability import get_logger

logger = get_logger("app.services.documents")

# Map a MIME type / extension hint to an extractor. Kept small and explicit.
_PDF_TYPES = {"application/pdf"}
_XLSX_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}
_DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}


def extract_text(data: bytes, *, mime_type: str | None, filename: str | None = None) -> str:
    """Extract plain text from document ``data``. Returns "" on any failure."""
    kind = _detect_kind(mime_type, filename)
    try:
        if kind == "pdf":
            return _extract_pdf(data)
        if kind == "xlsx":
            return _extract_xlsx(data)
        if kind == "docx":
            return _extract_docx(data)
    except Exception as exc:  # noqa: BLE001 - a bad file must not break ingestion
        logger.warning(
            "document_extraction_failed", extra={"kind": kind, "error": type(exc).__name__}
        )
        return ""
    return ""


def _detect_kind(mime_type: str | None, filename: str | None) -> str | None:
    mt = (mime_type or "").lower()
    if mt in _PDF_TYPES:
        return "pdf"
    if mt in _XLSX_TYPES:
        return "xlsx"
    if mt in _DOCX_TYPES:
        return "docx"
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return "pdf"
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if name.endswith(".docx"):
        return "docx"
    return None


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(parts).strip()


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    # read_only + data_only: fast, and returns computed values not formulas.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    workbook.close()
    return "\n".join(lines).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs if p.text).strip()
